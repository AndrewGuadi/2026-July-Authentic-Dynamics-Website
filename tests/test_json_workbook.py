import io
import json

import pytest
from openpyxl import load_workbook

from authentic_dynamics import create_app
from authentic_dynamics.blueprints.tools import json_tables
from authentic_dynamics.blueprints.tools.converters import ConversionError, convert_json

ORDERS = [
    {"id": "001", "customer": {"name": "Maya"},
     "items": [{"sku": "A1", "details": {"price": 12}, "tags": ["red", "small"]},
               {"sku": "B4", "quantity": 1}]},
    {"id": "002", "customer": {"name": "Alex"}, "items": [{"sku": "C1"}]},
]


def convert(value, **options):
    return convert_json(json.dumps(value).encode(), "xlsx", **options)


def workbook_rows(data):
    with io.BytesIO(data) as stream:
        book = load_workbook(stream)
        result = {}
        for sheet in book:
            values = list(sheet.values)
            result[sheet.title] = [dict(zip(values[0], row)) for row in values[1:]]
            assert sheet.freeze_panes == "A2"
            assert sheet.auto_filter.ref == sheet.dimensions
        book.close()
    return result


@pytest.mark.parametrize("mode,count,customer", [
    ("keep", 1, "customer"), ("flatten", 1, "customer.name"),
    ("related", 3, "customer"), ("combined", 3, "customer.name"),
])
def test_modes(mode, count, customer):
    rows = workbook_rows(convert(ORDERS, nesting=mode))
    assert len(rows) == count
    assert customer in rows["Records"][0]
    assert rows["Records"][0]["id"] == "001"
    if mode in {"keep", "flatten"}:
        assert json.loads(rows["Records"][0]["items"])[0]["sku"] == "A1"
        assert "@record_id" not in rows["Records"][0]


def test_related_links_order_and_grandchildren():
    rows = workbook_rows(convert(ORDERS))
    assert [row["@record_id"] for row in rows["Records"]] == ["1", "2"]
    items = rows["Records_items"]
    assert [row["@parent_id"] for row in items] == ["1", "1", "2"]
    assert [row["@index"] for row in items] == ["0", "1", "0"]
    assert [row["sku"] for row in items] == ["A1", "B4", "C1"]
    assert items[0]["details.price"] == "12"
    tags = rows["Records_items_tags"]
    assert [row["@parent_id"] for row in tags] == [items[0]["@record_id"]] * 2
    assert [row["@value"] for row in tags] == ["red", "small"]


def test_arrays_inside_objects_and_sibling_arrays_do_not_multiply_rows():
    value = {"customer": {"phones": ["123", "456"]}, "tags": [1, 2, 3]}
    rows = workbook_rows(convert(value))
    assert len(rows["Records"]) == 1
    assert len(rows["Records_customer.phones"]) == 2
    assert len(rows["Records_tags"]) == 3
    related = workbook_rows(convert(value, nesting="related"))
    assert json.loads(related["Records"][0]["customer"])["phones"] == ["123", "456"]


def test_mixed_empty_null_and_sparse_values():
    value = [{"items": [None, True, 3, "001", {"a": 1}, {}, [1, 2]], "empty": []},
             {"items": None, "object": {}}, {"other": "x"}]
    rows = workbook_rows(convert(value))
    assert rows["Records"][0]["empty"] == "[]"
    assert rows["Records_empty"] == []
    assert rows["Records"][1]["object"] == "{}"
    assert rows["Records"][1]["items"] is None
    items = rows["Records_items"]
    assert [row["@type"] for row in items] == [
        "null", "boolean", "number", "string", "object", "object", "array"]
    assert items[3]["@value"] == "001"
    assert items[4]["a"] == "1"
    assert items[5]["@type"] == "object"
    assert items[6]["@value"] == "[1, 2]"
    preview = convert(value, preview=True)
    assert any("Mixed arrays" in message for message in preview["warnings"])
    assert any("different fields" in message for message in preview["warnings"])


def test_depth_preserves_deeper_values():
    value = {"a": {"b": {"c": 1}}, "items": [{"more": [1, 2]}]}
    rows = workbook_rows(convert(value, max_depth=1))
    assert rows["Records"][0]["a.b"] == '{"c": 1}'
    assert rows["Records_items"][0]["more"] == '[1, 2]'
    preview = convert(value, max_depth=1, preview=True)
    assert any("deeper" in message for message in preview["warnings"])


def test_path_and_generated_column_collisions_and_formula_protection():
    value = {"a.b": "literal", "a": {"b": "nested"}, "@record_id": "original",
             "\\e": "slash", "": "empty", "=header": "=1+1",
             "items": [{"@parent_id": "original parent", "@value": "original value"}]}
    data = convert(value)
    rows = workbook_rows(data)
    main = rows["Records"][0]
    assert main["a\\.b"] == "literal" and main["a.b"] == "nested"
    assert main["\\@record_id"] == "original" and main["@record_id"] == "1"
    assert main["\\\\e"] == "slash" and main["\\e"] == "empty"
    assert rows["Records_items"][0]["\\@parent_id"] == "original parent"
    book = load_workbook(io.BytesIO(data))
    assert all(cell.data_type != "f" for sheet in book for row in sheet for cell in row)
    book.close()


def test_sheet_names_are_valid_unique_and_preview_matches_download():
    value = {"bad/name": [1], "bad:name": [2], "CASE": [3], "case": [4],
             "long" * 20: [5], "long" * 21: [6], "<script>": [7]}
    preview = convert(value, preview=True)
    book = load_workbook(io.BytesIO(convert(value)))
    assert len({name.casefold() for name in book.sheetnames}) == len(book.sheetnames)
    for sheet, info in zip(book, preview["sheets"]):
        assert sheet.title == info["name"]
        assert len(sheet.title) <= 31
        assert not set(sheet.title).intersection('\\/*?:[]')
        assert list(next(sheet.values)) == info["columns"]
        assert sheet.max_row - 1 == info["row_count"]
    book.close()


@pytest.mark.parametrize("options", [{"nesting": "bad"}, {"max_depth": 0},
                                     {"max_depth": 11}, {"max_depth": True}])
def test_invalid_options(options):
    with pytest.raises(ConversionError):
        convert({"a": 1}, **options)


@pytest.mark.parametrize("limit,value", [("MAX_ROWS", {"a": [1, 2, 3]}),
                                        ("MAX_SHEETS", {"a": [], "b": [], "c": []}),
                                        ("MAX_COLUMNS", {"a": 1, "b": 2, "c": 3}),
                                        ("MAX_CELLS", {"a": 1, "b": 2})])
def test_workbook_limits_include_generated_fields(monkeypatch, limit, value):
    monkeypatch.setattr(json_tables, limit, 3)
    with pytest.raises(ConversionError):
        convert(value, preview=True)


def test_late_columns_count_existing_empty_cells(monkeypatch):
    monkeypatch.setattr(json_tables, "MAX_CELLS", 8)
    with pytest.raises(ConversionError, match="200,000"):
        convert([{"a": 1}, {"b": 2}, {"c": 3}], nesting="flatten", preview=True)


def test_non_excel_ignores_nesting_and_preview_rejects_other_formats():
    value = b'{"a":{"b":1},"items":[1,2]}'
    assert convert_json(value, "csv", "combined") == convert_json(value, "csv", "keep")
    assert json.loads(convert_json(value, "txt")) == json.loads(value)
    with pytest.raises(ConversionError, match="Excel"):
        convert_json(value, "csv", preview=True)


def test_preview_route_upload_text_html_and_csrf():
    app = create_app({"TESTING": True, "SECRET_KEY": "test",
                      "SQLALCHEMY_DATABASE_URI": "sqlite://"})
    client = app.test_client()
    assert client.post('/tools/json-converter', data={'action': 'preview'}).status_code == 400
    app.config['WTF_CSRF_ENABLED'] = False
    for source in [{"json_text": json.dumps(ORDERS)},
                   {"file": (io.BytesIO(json.dumps(ORDERS).encode()), "orders.json")}]:
        response = client.post('/tools/json-converter', data={
            **source, "format": "xlsx", "action": "preview"}, headers={"Accept": "application/json"})
        assert response.status_code == 200
        assert response.headers["Cache-Control"] == "no-store"
        assert "Content-Disposition" not in response.headers
        assert len(response.json["sheets"]) == 3
    response = client.post('/tools/json-converter', data={
        "json_text": '{"<script>":[1]}', "format": "xlsx", "action": "preview",
        "nesting": "combined", "max_depth": "2"})
    assert response.status_code == 200
    assert b'Workbook preview' in response.data and b'&lt;script&gt;' in response.data
    assert b'<script>\xc2\xb7' not in response.data
    assert b'value="2" selected' in response.data
    for options in [{"max_depth": "garbage"}, {"nesting": "garbage"}]:
        response = client.post('/tools/json-converter', data={
            "json_text": '{"a":1}', "action": "preview", **options},
            headers={"Accept": "application/json"})
        assert response.status_code == 400 and response.json["error"]
