"""Bounded normalization shared by XLSX downloads and workbook previews."""

import io
import json
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .conversion_errors import ConversionError

MODES = {"keep", "flatten", "related", "combined"}
MAX_SHEETS = 20
MAX_ROWS = 50000
MAX_COLUMNS = 100
MAX_CELLS = 200000


def cell_text(value):
    text = value if isinstance(value, str) else (
        "" if value is None else json.dumps(value, ensure_ascii=False, allow_nan=False))
    if len(text) > 32767:
        raise ConversionError("Keep each cell under 32,768 characters. Try flattening or related sheets.")
    if re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]", text):
        raise ConversionError("Table cells cannot contain unsupported control characters.")
    return text


def field_name(key):
    # Escape literal dots/backslashes and reserve @ names for generated metadata.
    name = key.replace("\\", "\\\\").replace(".", "\\.") or "\\e"
    return "\\" + name if name.startswith("@") else name


def normalize_workbook(records, mode="combined", max_depth=5):
    if mode not in MODES:
        raise ConversionError("Choose a supported nested data option.")
    if type(max_depth) is not int or not 1 <= max_depth <= 10:
        raise ConversionError("Choose a nesting depth from 1 to 10.")
    flatten = mode in {"flatten", "combined"}
    related = mode in {"related", "combined"}
    sheets = {}
    warnings = set()
    total_rows = total_cells = 0

    def add_header(sheet, key):
        nonlocal total_cells
        if key in sheet["columns"]:
            return
        if len(sheet["columns"]) >= MAX_COLUMNS:
            raise ConversionError("Use at most 100 columns per worksheet. Try a shallower depth.")
        cell_text(key)
        total_cells += len(sheet["rows"]) + 1
        if total_cells > MAX_CELLS:
            raise ConversionError("Use at most 200,000 cells across all worksheets, including headers.")
        sheet["columns"].append(key)

    def get_sheet(path, parent=None, field=None):
        if path in sheets:
            return sheets[path]
        if len(sheets) >= MAX_SHEETS:
            raise ConversionError("Use at most 20 worksheets. Try a shallower depth or keep arrays in cells.")
        base = "Records" if parent is None else f"{parent['name']}_{field}"
        base = re.sub(r"[\\/*?:\[\]\x00-\x1f]", "_", base)[:31].strip("'") or "Sheet"
        name, suffix = base, 2
        used = {sheet["name"].casefold() for sheet in sheets.values()}
        while name.casefold() in used:
            tail = f"_{suffix}"
            name = base[:31 - len(tail)] + tail
            suffix += 1
        sheet = {"name": name, "columns": [], "rows": [],
                 "parent_sheet": parent["name"] if parent else None, "source_field": field}
        sheets[path] = sheet
        if related:
            for key in (["@record_id", "@parent_id", "@index", "@type", "@value"]
                        if parent else ["@record_id"]):
                add_header(sheet, key)
        return sheet

    def add_record(sheet, path, item, depth=1, parent_id=None, index=None):
        nonlocal total_rows, total_cells
        total_rows += 1
        if total_rows > MAX_ROWS:
            raise ConversionError("Use at most 50,000 data rows across all worksheets.")
        row = {}
        record_id = str(len(sheet["rows"]) + 1)
        if related:
            row["@record_id"] = record_id
        if parent_id is not None:
            kind = ("object" if isinstance(item, dict) else "array" if isinstance(item, list)
                    else "null" if item is None else "boolean" if isinstance(item, bool)
                    else "number" if isinstance(item, (int, float)) else "string")
            row.update({"@parent_id": parent_id, "@index": str(index), "@type": kind})
            if not isinstance(item, dict):
                row["@value"] = cell_text(item)
                if isinstance(item, list):
                    warnings.add("Arrays inside array items stay as JSON in the @value column.")

        def visit(value, keys, level):
            label = ".".join(field_name(key) for key in keys) if mode != "keep" else keys[0]
            expandable = (flatten and isinstance(value, dict) and value) or (
                related and isinstance(value, list))
            if expandable and level > max_depth:
                warnings.add("Values deeper than the selected depth stay as JSON text in cells.")
            elif flatten and isinstance(value, dict) and value:
                for key, child in value.items():
                    visit(child, (*keys, key), level + 1)
                return
            elif related and isinstance(value, list):
                child_path = (*path, keys)
                child_sheet = get_sheet(child_path, sheet, label)
                row[label] = f"[{len(value)} items → {child_sheet['name']}]" if value else "[]"
                if len({type(child) for child in value}) > 1:
                    warnings.add("Mixed arrays use @type and @value alongside object fields on child sheets.")
                for child_index, child in enumerate(value):
                    add_record(child_sheet, child_path, child, level + 1, record_id, child_index)
                return
            row[label] = cell_text(value)

        if isinstance(item, dict):
            for key, value in item.items():
                visit(value, (key,), depth)
            shape = {key for key in row if not related or not key.startswith("@")}
            if "shape" in sheet and shape != sheet["shape"]:
                warnings.add("Objects with different fields share a combined set of columns; missing fields are blank.")
            sheet.setdefault("shape", shape)
        for key in row:
            add_header(sheet, key)
        total_cells += len(sheet["columns"])
        if total_cells > MAX_CELLS:
            raise ConversionError("Use at most 200,000 cells across all worksheets, including headers.")
        sheet["rows"].append(row)

    main = get_sheet(())
    for record in records:
        add_record(main, (), record)
    if not main["columns"] or (related and main["columns"] == ["@record_id"]):
        raise ConversionError("Use at least one field for a table.")
    if mode == "related":
        warnings.add("Objects stay in cells; only arrays directly on each row become related sheets.")
    return {"sheets": list(sheets.values()), "warnings": sorted(warnings)}


def workbook_preview(workbook):
    return {"sheets": [{"name": sheet["name"], "columns": sheet["columns"],
                        "row_count": len(sheet["rows"]), "parent_sheet": sheet["parent_sheet"],
                        "source_field": sheet["source_field"]} for sheet in workbook["sheets"]],
            "warnings": workbook["warnings"]}


def write_workbook(normalized):
    book = Workbook()
    book.remove(book.active)
    for table in normalized["sheets"]:
        sheet = book.create_sheet(table["name"])
        headers = table["columns"]
        for row_index, row in enumerate([dict(zip(headers, headers)), *table["rows"]], start=1):
            for column_index, header in enumerate(headers, start=1):
                cell = sheet.cell(row_index, column_index, row.get(header, ""))
                cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if row_index == 1:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="26371B")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for index, header in enumerate(headers, start=1):
            width = max([len(header), *(len(row.get(header, "")) for row in table["rows"][:50])])
            sheet.column_dimensions[get_column_letter(index)].width = min(45, max(14, width + 2))
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()
