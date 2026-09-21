"""Bounded file conversions. Uploaded content is never persisted by the app."""

import csv
import io
import json
import math
import re
import zipfile
from xml.etree import ElementTree as ET

import pypdfium2 as pdfium
from openpyxl import Workbook

from .conversion_errors import ConversionError
from .json_tables import cell_text, normalize_workbook, workbook_preview, write_workbook
from .pdf_tables import PDF_LOCK, build_pdf

MAX_FILE_BYTES = 10 * 1024 * 1024


def convert_json(data, output_format, nesting="combined", max_depth=5, preview=False,
                 pdf_options=None):
    if output_format not in {"xlsx", "csv", "tsv", "json", "txt", "docx", "pdf"}:
        raise ConversionError("Choose a supported output format.")
    if preview and output_format not in {"xlsx", "pdf"}:
        raise ConversionError("Choose Excel or PDF to preview the output.")
    if not data or len(data) > MAX_FILE_BYTES:
        raise ConversionError("Provide nonempty JSON no larger than 10 MiB.")

    def unique_object(pairs):
        result = {}
        for key, value in pairs:
            if key in result:
                raise ConversionError("JSON contains duplicate object keys. Make each key unique.")
            result[key] = value
        return result

    def invalid_constant(value):
        raise ConversionError("JSON cannot contain NaN or Infinity.")

    try:
        value = json.loads(data.decode("utf-8-sig"), object_pairs_hook=unique_object,
                           parse_constant=invalid_constant)
        # Reject numeric overflow and escaped surrogate characters as well.
        formatted = json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False).encode()
    except json.JSONDecodeError as exc:
        raise ConversionError(f"Invalid JSON at line {exc.lineno}, column {exc.colno}: {exc.msg}.") from exc
    except (UnicodeError, RecursionError, ValueError) as exc:
        if isinstance(exc, ConversionError):
            raise
        raise ConversionError("Use valid UTF-8 JSON with finite numbers and less nesting.") from exc
    if output_format in {"json", "txt"}:
        return formatted
    records = [value] if isinstance(value, dict) else value
    if not isinstance(records, list) or not records or not all(isinstance(row, dict) for row in records):
        raise ConversionError("For tables, use an object or a nonempty array of objects. "
                              "Formatted JSON and plain text support any JSON value.")
    if len(records) > 50000:
        raise ConversionError("Use at most 50,000 data rows.")
    if output_format == "xlsx":
        normalized = normalize_workbook(records, nesting, max_depth)
        return workbook_preview(normalized) if preview else write_workbook(normalized)
    if output_format == "pdf":
        result, summary, _ = build_pdf(records, **(pdf_options or {}))
        return {"pdf": result, "summary": summary} if preview else result
    headers = list(dict.fromkeys(key for row in records for key in row))
    if not headers or len(headers) > 100:
        raise ConversionError("Use 1–100 columns for a spreadsheet.")
    if (len(records) + 1) * len(headers) > 200000:
        raise ConversionError("Use at most 200,000 cells per file.")

    rows = [[cell_text(key) for key in headers]]
    rows.extend([cell_text(row.get(key)) for key in headers] for row in records)
    if output_format == "docx":
        return _json_docx(rows)
    if output_format in {"csv", "tsv"}:
        output = io.StringIO(newline="")
        writer = csv.writer(output, delimiter="," if output_format == "csv" else "\t")
        writer.writerows([["'" + cell if cell.lstrip().startswith(("=", "+", "-", "@"))
                           else cell for cell in row] for row in rows])
        return output.getvalue().encode("utf-8-sig")


def _json_docx(rows):
    """Create a minimal, readable Word document without a third-party dependency."""
    body = []
    for row in rows:
        cells = "".join(f"<w:tc><w:tcPr/><w:p><w:r><w:t xml:space=\"preserve\">"
                        f"{_xml_escape(value)}</w:t></w:r></w:p></w:tc>" for value in row)
        body.append(f"<w:tr>{cells}</w:tr>")
    document = ("<?xml version=\"1.0\" encoding=\"UTF-8\" standalone=\"yes\"?>"
                "<w:document xmlns:w=\"http://schemas.openxmlformats.org/wordprocessingml/2006/main\">"
                "<w:body><w:tbl><w:tblPr><w:tblBorders><w:top w:val=\"single\"/>"
                "<w:left w:val=\"single\"/><w:bottom w:val=\"single\"/><w:right w:val=\"single\"/>"
                "<w:insideH w:val=\"single\"/><w:insideV w:val=\"single\"/></w:tblBorders></w:tblPr>"
                f"{''.join(body)}</w:tbl><w:sectPr/></w:body></w:document>")
    content_types = ("<?xml version=\"1.0\" encoding=\"UTF-8\"?>"
                     "<Types xmlns=\"http://schemas.openxmlformats.org/package/2006/content-types\">"
                     "<Default Extension=\"rels\" ContentType=\"application/vnd.openxmlformats-package.relationships+xml\"/>"
                     "<Default Extension=\"xml\" ContentType=\"application/xml\"/>"
                     "<Override PartName=\"/word/document.xml\" ContentType=\"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml\"/>"
                     "</Types>")
    rels = ("<?xml version=\"1.0\" encoding=\"UTF-8\"?>"
            "<Relationships xmlns=\"http://schemas.openxmlformats.org/package/2006/relationships\">"
            "<Relationship Id=\"rId1\" Type=\"http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument\" Target=\"word/document.xml\"/>"
            "</Relationships>")
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", rels)
        archive.writestr("word/document.xml", document)
    return output.getvalue()


def _xml_escape(value):
    return (value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            .replace('"', "&quot;").replace("'", "&apos;"))


def convert_pdf(data, output_format, dpi):
    if output_format not in {"png", "webp", "jpeg"} or dpi not in {72, 150, 300}:
        raise ConversionError("Choose a supported image format and resolution.")
    if not data.startswith(b"%PDF-"):
        raise ConversionError("Choose a valid PDF file.")
    with PDF_LOCK:
        try:
            document = pdfium.PdfDocument(data)
        except pdfium.PdfiumError as exc:
            raise ConversionError("This PDF could not be opened. Check it is valid and unlocked.") from exc
        with document:
            if not 1 <= len(document) <= 40:
                raise ConversionError("Choose a PDF with 1–40 pages.")
            scale = dpi / 72
            total = 0
            for index in range(len(document)):
                width, height = document.get_page_size(index)
                if not all(math.isfinite(n) and n > 0 for n in (width, height)):
                    raise ConversionError("This PDF contains an invalid page size.")
                pixels = math.ceil(width * scale) * math.ceil(height * scale)
                total += pixels
                if (pixels > 20_000_000 or total > 80_000_000
                        or max(width, height) * scale > 16000):
                    raise ConversionError("This PDF is too large to render. Try fewer pages or a lower resolution.")
            output = io.BytesIO()
            archive = zipfile.ZipFile(output, "w") if len(document) > 1 else None
            try:
                for index in range(len(document)):
                    page = document[index]
                    try:
                        bitmap = page.render(scale=scale)
                        try:
                            with bitmap.to_pil().convert("RGB") as image:
                                page_output = io.BytesIO()
                                image.save(page_output, format=output_format.upper(), quality=90)
                        finally:
                            bitmap.close()
                    finally:
                        page.close()
                    if archive:
                        archive.writestr(f"page-{index + 1:03d}.{output_format}", page_output.getvalue())
                    else:
                        output.write(page_output.getvalue())
                    if output.tell() > 100 * 1024 * 1024:
                        raise ConversionError("The result is too large. Try fewer pages or a lower resolution.")
            except pdfium.PdfiumError as exc:
                raise ConversionError("A page could not be rendered. Try exporting the PDF again.") from exc
            finally:
                if archive:
                    archive.close()
            return output.getvalue(), "zip" if archive else output_format


def convert_csv(data, output_format, delimiter, has_headers):
    if output_format not in {"xlsx", "json", "xml", "tsv"}:
        raise ConversionError("Choose a supported output format.")
    if delimiter not in {",", ";", "\t", "|"}:
        raise ConversionError("Choose a supported delimiter.")
    try:
        content = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ConversionError("Save your CSV with UTF-8 encoding and try again.") from exc
    if re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]", content):
        raise ConversionError("The CSV contains unsupported control characters.")
    rows = []
    try:
        for row in csv.reader(io.StringIO(content, newline=""), delimiter=delimiter, strict=True):
            if not row:
                continue
            if len(row) > 100 or len(rows) >= 50001:
                raise ConversionError("Use at most 50,000 data rows and 100 columns.")
            if rows and len(row) != len(rows[0]):
                raise ConversionError(f"Row {len(rows) + 1} has a different number of columns. Check your delimiter and CSV.")
            if any(len(value) > 32767 for value in row):
                raise ConversionError("Keep each cell under 32,768 characters.")
            rows.append(row)
            if len(rows) * len(row) > 200_000:
                raise ConversionError("Use at most 200,000 cells per file.")
    except csv.Error as exc:
        raise ConversionError("The CSV could not be read. Check quoting and delimiters.") from exc
    if not rows:
        raise ConversionError("The CSV is empty. Choose a file with data.")
    if not has_headers and len(rows) > 50000:
        raise ConversionError("Use at most 50,000 data rows.")
    headers = rows[0] if has_headers else [f"column_{i + 1}" for i in range(len(rows[0]))]
    records = rows[1:] if has_headers else rows
    if any(not header.strip() for header in headers) or len(set(headers)) != len(headers):
        raise ConversionError("Column headers must be nonempty and unique, or turn off the header-row option.")
    if output_format == "json":
        return json.dumps([dict(zip(headers, row)) for row in records], ensure_ascii=False, indent=2).encode()
    if output_format == "xml":
        root = ET.Element("rows")
        for row in records:
            element = ET.SubElement(root, "row")
            for header, value in zip(headers, row):
                ET.SubElement(element, "field", name=header).text = value
        return ET.tostring(root, encoding="utf-8", xml_declaration=True)
    if output_format == "tsv":
        output = io.StringIO(newline="")
        writer = csv.writer(output, delimiter="\t")
        # Neutralize spreadsheet formulas when a TSV is opened in Excel.
        writer.writerows([
            ["'" + value if value.lstrip().startswith(("=", "+", "-", "@")) else value for value in row]
            for row in [headers, *records]
        ])
        return output.getvalue().encode("utf-8-sig")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Converted CSV"
    for row_index, row in enumerate([headers, *records], start=1):
        for column_index, value in enumerate(row, start=1):
            cell = sheet.cell(row_index, column_index, value)
            cell.data_type = "s"  # Preserve leading zeros and never execute uploaded formulas.
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
